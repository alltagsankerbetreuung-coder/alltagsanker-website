#!/usr/bin/env python3
"""Erstellt die Word-Vorlage für Briefe (DIN 5008, Form B):
druckvorlagen/word/alltagsfels-briefvorlage.docx

Aufruf (aus dem Hauptordner):  python3 druckvorlagen/word-vorlage-erstellen.py
Braucht das Python-Paket python-docx."""
import copy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

HIER = Path(__file__).parent
LOGO = HIER / "logo" / "alltagsfels-logo.png"
ZIEL = HIER / "word" / "alltagsfels-briefvorlage.docx"
LINIE = HIER / "word" / "_falzmarke.png"

SCHRIFT = "Atkinson Hyperlegible"
TEXT = RGBColor(0x26, 0x22, 0x1E)
BLAU = RGBColor(0x17, 0x3A, 0x66)
GRAU = RGBColor(0x5A, 0x52, 0x4A)
EMU_MM = 36000


def schrift(run, groesse=None, fett=None, farbe=None, gelb=False):
    run.font.name = SCHRIFT
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), SCHRIFT)
    if groesse:
        run.font.size = Pt(groesse)
    if fett is not None:
        run.font.bold = fett
    if farbe is not None:
        run.font.color.rgb = farbe
    if gelb:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return run


def absatz(container, teile=(), groesse=12, vor=0, nach=0, zeile=None):
    """teile: Liste aus (Text, Optionen-dict)"""
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(vor)
    pf.space_after = Pt(nach)
    if zeile:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(zeile)
    for text, opt in teile:
        schrift(p.add_run(text), groesse=opt.get("g", groesse), fett=opt.get("b"),
                farbe=opt.get("f"), gelb=opt.get("gelb", False))
    return p


def verankern(run, x_mm, y_mm):
    """Macht aus einem eingefügten Bild ein frei auf der Seite platziertes Bild."""
    inline = run._r.find(".//" + qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    docpr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = parse_xml(
        f'<wp:anchor {nsdecls("wp", "a", "pic", "r")} distT="0" distB="0" distL="0" distR="0" '
        f'simplePos="0" relativeHeight="251659264" behindDoc="1" locked="1" layoutInCell="1" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{int(x_mm * EMU_MM)}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{int(y_mm * EMU_MM)}</wp:posOffset></wp:positionV>'
        f"</wp:anchor>"
    )
    anchor.append(copy.deepcopy(extent))
    anchor.append(parse_xml(f'<wp:effectExtent {nsdecls("wp")} l="0" t="0" r="0" b="0"/>'))
    anchor.append(parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>'))
    anchor.append(copy.deepcopy(docpr))
    anchor.append(parse_xml(f'<wp:cNvGraphicFramePr {nsdecls("wp")}/>'))
    anchor.append(copy.deepcopy(graphic))
    inline.getparent().replace(inline, anchor)


def zellen_ohne_rand(tabelle, oben_linie=False):
    tbl = tabelle._tbl
    tblpr = tbl.tblPr
    linie = '<w:top w:val="single" w:sz="6" w:space="0" w:color="E9D8C1"/>' if oben_linie else '<w:top w:val="nil"/>'
    tblpr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}>{linie}<w:left w:val="nil"/><w:bottom w:val="nil"/>'
        f'<w:right w:val="nil"/><w:insideH w:val="nil"/><w:insideV w:val="nil"/></w:tblBorders>'))
    tblpr.append(parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:left w:w="0" w:type="dxa"/>'
        f'<w:bottom w:w="0" w:type="dxa"/><w:right w:w="0" w:type="dxa"/></w:tblCellMar>'))
    tblpr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))


def spaltenbreiten(tabelle, breiten_mm):
    tabelle.autofit = False
    grid = tabelle._tbl.tblGrid
    for gc, b in zip(grid.findall(qn("w:gridCol")), breiten_mm):
        gc.set(qn("w:w"), str(int(b / 25.4 * 1440)))
    for zeile in tabelle.rows:
        for zelle, b in zip(zeile.cells, breiten_mm):
            zelle.width = Mm(b)


def zelle_leeren(zelle):
    for p in list(zelle.paragraphs):
        p._p.getparent().remove(p._p)


def main():
    # kleine graue Linie für die Falz- und Lochmarken
    Image.new("RGB", (60, 3), (154, 144, 135)).save(LINIE)

    doc = Document()
    doc.core_properties.title = "Alltagsfels – Briefvorlage"
    doc.core_properties.author = "Alltagsfels"

    normal = doc.styles["Normal"]
    normal.font.name = SCHRIFT
    normal.font.size = Pt(12)
    normal.font.color.rgb = TEXT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SCHRIFT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15

    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin = Mm(25), Mm(20)
    s.top_margin, s.bottom_margin = Mm(45), Mm(34)
    s.header_distance, s.footer_distance = Mm(10), Mm(8)

    # ---------- Kopfzeile: Logo und Falzmarken (frei platziert) ----------
    kopf = s.header.paragraphs[0]
    logo_run = kopf.add_run()
    logo_run.add_picture(str(LOGO), width=Mm(72))
    verankern(logo_run, 190 - 72, 13)
    for y in (105, 148.5, 210):
        r = kopf.add_run()
        r.add_picture(str(LINIE), width=Mm(7 if y == 148.5 else 5), height=Mm(0.3))
        verankern(r, 0, y)

    # ---------- Anschriftfeld und Informationsblock ----------
    t = doc.add_table(rows=2, cols=3)
    zellen_ohne_rand(t)
    spaltenbreiten(t, [80, 20, 65])
    t.rows[0].height, t.rows[0].height_rule = Mm(17.7), WD_ROW_HEIGHT_RULE.EXACTLY
    t.rows[1].height, t.rows[1].height_rule = Mm(27.3), WD_ROW_HEIGHT_RULE.EXACTLY

    # Rücksendeangabe
    z = t.cell(0, 0)
    zelle_leeren(z)
    z.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    p = absatz(z, [("Alltagsfels UG (haftungsbeschränkt) · Batenbrockstr. 66 · 46238 Bottrop", {"f": GRAU})],
               groesse=7, nach=2)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="1" w:color="9A9087"/></w:pBdr>'))

    # Empfänger
    z = t.cell(1, 0)
    zelle_leeren(z)
    for zeile in ("Frau / Herrn", "Vorname Nachname", "Straße Hausnummer", "PLZ Ort"):
        absatz(z, [(zeile, {})], zeile=14)

    # Informationsblock (ab 50 mm von oben)
    info = t.cell(0, 2).merge(t.cell(1, 2))
    zelle_leeren(info)
    absatz(info, [("Ihr Ansprechpartner: ", {"b": True, "f": BLAU}), ("Julian Wehner", {})], groesse=11, vor=14)
    absatz(info, [("Telefon: ", {"b": True, "f": BLAU}), ("[TELEFON]", {"gelb": True, "b": True})], groesse=11)
    absatz(info, [("E-Mail: ", {"b": True, "f": BLAU}), ("[E-MAIL]", {"gelb": True, "b": True})], groesse=11)
    absatz(info, [("Internet: ", {"b": True, "f": BLAU}), ("www.alltagsfels.de", {})], groesse=11)
    absatz(info, [("Datum: ", {"b": True, "f": BLAU}), ("TT.MM.JJJJ", {})], groesse=11, vor=8)

    # ---------- Betreff und Brieftext ----------
    absatz(doc, [("Betreff Ihres Briefes", {"b": True})], vor=24)
    absatz(doc, [], vor=12)
    absatz(doc, [("Sehr geehrte Damen und Herren,", {})], vor=12, nach=12)
    absatz(doc, [("hier schreiben Sie Ihren Brief. Schreiben Sie einfach über diesen Text. "
                  "Die Schrift ist groß und gut lesbar. Absätze trennen Sie mit einer Leerzeile.", {})], nach=12)
    absatz(doc, [("Mit freundlichen Grüßen", {})], nach=36)
    absatz(doc, [("Julian Wehner", {})])
    absatz(doc, [("Geschäftsführer", {})])

    # ---------- Fußzeile ----------
    fuss = s.footer
    zelle_leeren(fuss)
    ft = fuss.add_table(rows=1, cols=4, width=Mm(165))
    zellen_ohne_rand(ft, oben_linie=True)
    spaltenbreiten(ft, [50, 45, 40, 30])
    spalten = [
        [("Alltagsfels UG", {"b": True, "f": BLAU}), ("(haftungsbeschränkt)", {"b": True, "f": BLAU}),
         ("Batenbrockstraße 66", {}), ("46238 Bottrop", {})],
        [("Geschäftsführer: Julian Wehner", {}), ("Amtsgericht Gelsenkirchen", {}), ("HRB 20165", {})],
        [("Telefon ", {}, "[TELEFON]"), ("E-Mail ", {}, "[E-MAIL]"), ("www.alltagsfels.de", {})],
        [("Bankverbindung", {}), ("IBAN ", {}, "[IBAN]")],
    ]
    for zelle, zeilen in zip(ft.rows[0].cells, spalten):
        zelle_leeren(zelle)
        for i, eintrag in enumerate(zeilen):
            teile = [(eintrag[0], dict(eintrag[1], f=eintrag[1].get("f", GRAU)))]
            if len(eintrag) == 3:
                teile.append((eintrag[2], {"gelb": True, "b": True, "f": RGBColor(0, 0, 0)}))
            absatz(zelle, teile, groesse=8.5, vor=5 if i == 0 else 0, zeile=11)

    doc.save(ZIEL)
    LINIE.unlink()
    print("fertig:", ZIEL)


if __name__ == "__main__":
    main()
